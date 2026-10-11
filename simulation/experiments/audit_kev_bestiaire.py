"""Audit : « des nœuds proches de la matière noire à l'échelle du keV » (bestiaire).

1. Ce que dit TQNT_catalog_V6_complete.xlsx (Q_désint en keV, τ longs) — et sur
   quel barème (commit 809d14d : masses Ω_ts V5, archivées).
2. Les mêmes canaux sous la loi géométrique canonique (ACPR) : ouverts ou fermés ?
3. Les structures keV qui subsistent sous ACPR entre états sombres (Q_phys=0).
4. L'époque cosmique de chaque τ, comparée à la fenêtre du défaut F (T = 1-10 keV).
5. Ce qu'un canal keV ferait physiquement : décroissance = coup de pied (chauffe),
   refroidissement = collisions (taux à T ~ keV pour chaque espèce du bestiaire).
"""
import sys, math
sys.path.insert(0, '.')
import openpyxl
from theory.ropelength_acpr import ROPELENGTH_ACPR, m_topo

def f(x):
    try: return float(str(x).replace(',', '.'))
    except Exception: return None

def comps(cfg): return [c.strip() for c in cfg.split('+')]

def T_of_t(t):          # ère de radiation, g* = 3,36 : t = 1,31 s (1 MeV/T)^2
    return 1000.0 / math.sqrt(t / 1.31)      # keV

def t_of_T(T_keV):
    return 1.31 * (1000.0 / T_keV) ** 2

print("=== 1. TQNT_catalog_V6_complete.xlsx : Q_désint <= 0,3 MeV (barème Ω_ts, commit 809d14d) ===")
ws = openpyxl.load_workbook('docs/TQNT_catalog_V6_complete.xlsx', read_only=True, data_only=True).active
rows = list(ws.iter_rows(values_only=True))[1:]
kev = []
for r in rows:
    Q = f(r[11]); tau = f(r[12])
    if Q is not None and 0 < Q <= 0.3 and tau:
        kev.append((r[1], r[10], Q, tau, f(r[4])))
kev.sort(key=lambda x: x[2])
print(f"{'config':>18} {'-> produit':>20} {'Q (keV)':>8} {'τ (s)':>9} {'T(τ)':>9}   canal sous ACPR")
for cfg, dau, Q, tau, m in kev:
    try:
        mA = sum(m_topo(k) for k in comps(cfg)); mB = sum(m_topo(k) for k in comps(dau))
        acpr = f"Q_ACPR = {mA - mB:+.0f} MeV -> {'OUVERT' if mA > mB else 'FERMÉ'}"
    except KeyError:
        acpr = "10_1 hors table ACPR"
    T = T_of_t(tau)
    Ts = f"{T:.0f} keV" if T >= 1 else f"{T*1000:.0f} eV"
    print(f"{cfg:>18} {dau:>20} {Q*1000:8.0f} {tau:9.1e} {Ts:>9}   {acpr}")

print("\n=== 2. Fenêtre du défaut F : M_DM causale = 3e5 Msun à T = 7 keV ===")
for T in (10, 7, 1):
    print(f"  T = {T:2d} keV  <->  t = {t_of_T(T):.1e} s")

print("\n=== 3. Sous ACPR : quasi-dégénérescences entre états sombres (Q_phys = 0) < 1 MeV ===")
ws2 = openpyxl.load_workbook('docs/TQNT_catalog_complete.xlsx', read_only=True, data_only=True)['Entrelacs complets (1539)']
dark = []
for i, r in enumerate(ws2.iter_rows(values_only=True)):
    if i < 2: continue
    if r[0] is None: break
    q = f(r[7])
    if q is None or abs(q) > 1e-9: continue
    try: m = sum(m_topo(k) for k in comps(r[1]))
    except KeyError: continue
    dark.append((m, r[1], r[3]))
dark.sort()
for a, b in zip(dark, dark[1:]):
    d = (b[0] - a[0]) * 1000
    if d < 1000:
        print(f"  {a[1]:>16} (c{a[2]}) {a[0]:8.1f}  |  {b[1]:>16} (c{b[2]}) {b[0]:8.1f}   Δ = {d:6.1f} keV  (compositions disjointes : pas un canal)")

print("\n=== 4. Un canal keV ferait quoi ? ===")
m_er = 1.8e3   # MeV
for Q in (0.03, 0.23):
    v = math.sqrt(2 * Q / m_er)
    print(f"  décroissance Q = {Q*1000:.0f} keV sur {m_er/1e3:.1f} GeV : coup de pied v = {v:.1e} c = {v*3e5:.0f} km/s  (CHAUFFE, ne refroidit pas)")
# refroidissement = collisions ; densité d'éremnions à T
rho0 = 2.3e-27      # kg/m3 (DM aujourd'hui)
m_kg = m_er * 1.783e-30
for T in (7.0, 1.0):
    z1 = T * 1e3 / 2.35e-4          # 1+z (T_CMB0 = 0,235 meV)
    n = rho0 * z1 ** 3 / m_kg * 1e-6  # cm^-3
    v = math.sqrt(3 * T * 1e-3 / m_er) * 3e10   # cm/s thermique
    t = t_of_T(T)
    print(f"  T = {T:.0f} keV : n_DM = {n:.1e} cm^-3, v_th = {v/3e10:.1e} c, t = {t:.1e} s")
    for nom, sig in (("réservoir c>=23 (σ_eq)", 1e-48), ("bande c=22", 1e-33), ("bande c=20", 1e-29), ("bande c=19", 1e-27)):
        print(f"      {nom:>24} : collisions par particule et par âge = {n*sig*v*t:.1e}")
print("  (les bandes c=19-22 sont couplées à la naissance et annihilent : reliques ~nulles ; σv(c=19)~3e-17 cm3/s -> Ω h² ~ 1e-10)")
